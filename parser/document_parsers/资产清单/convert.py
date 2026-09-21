from __future__ import annotations

import argparse
import json
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl
import xlrd


ASSET_FIELDS = (
    "资产IP",
    "映射IP",
    "资产类型",
    "资产型号",
    "设备厂商",
    "域名",
    "网络位置",
    "定级对象名称",
    "责任人",
)


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _normal_header(value: str) -> str:
    return re.sub(r"[\s\u3000]+", "", value).replace("：", "").replace(":", "")


HEADER_ALIASES = {_normal_header(field): field for field in ASSET_FIELDS}


@dataclass
class SheetData:
    title: str
    rows: list[list[str]]
    merged_ranges: list[tuple[int, int, int, int]]

    def __post_init__(self) -> None:
        self._merged_values: dict[tuple[int, int], str] = {}
        for row_start, row_end, col_start, col_end in self.merged_ranges:
            anchor = self.raw(row_start, col_start)
            for row_index in range(row_start, row_end):
                for col_index in range(col_start, col_end):
                    self._merged_values[(row_index, col_index)] = anchor

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def column_count(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def raw(self, row_index: int, col_index: int) -> str:
        if row_index < 0 or row_index >= len(self.rows):
            return ""
        row = self.rows[row_index]
        if col_index < 0 or col_index >= len(row):
            return ""
        return row[col_index]

    def value(self, row_index: int, col_index: int) -> str:
        raw = self.raw(row_index, col_index)
        if raw:
            return raw
        return self._merged_values.get((row_index, col_index), "")


def _read_xlsx(path: Path) -> list[SheetData]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    sheets: list[SheetData] = []
    for worksheet in workbook.worksheets:
        rows = [
            [_clean_text(worksheet.cell(row_index, col_index).value) for col_index in range(1, worksheet.max_column + 1)]
            for row_index in range(1, worksheet.max_row + 1)
        ]
        merged = [
            (item.min_row - 1, item.max_row, item.min_col - 1, item.max_col)
            for item in worksheet.merged_cells.ranges
        ]
        sheets.append(SheetData(worksheet.title, rows, merged))
    workbook.close()
    return sheets


def _xls_cell_text(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> str:
    if cell.ctype == xlrd.XL_CELL_DATE:
        return _clean_text(xlrd.xldate_as_datetime(cell.value, book.datemode))
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    return _clean_text(cell.value)


def _read_xls(path: Path) -> list[SheetData]:
    try:
        workbook = xlrd.open_workbook(path, formatting_info=True)
    except (NotImplementedError, xlrd.XLRDError):
        workbook = xlrd.open_workbook(path)
    sheets: list[SheetData] = []
    for worksheet in workbook.sheets():
        rows = [
            [_xls_cell_text(workbook, worksheet.cell(row_index, col_index)) for col_index in range(worksheet.ncols)]
            for row_index in range(worksheet.nrows)
        ]
        sheets.append(SheetData(worksheet.name, rows, [tuple(item) for item in worksheet.merged_cells]))
    return sheets


def read_workbook(path: str | Path) -> list[SheetData]:
    source = Path(path)
    if source.suffix.lower() == ".xlsx":
        return _read_xlsx(source)
    if source.suffix.lower() == ".xls":
        return _read_xls(source)
    raise ValueError(f"仅支持 .xls 和 .xlsx：{source}")


def _find_header(sheet: SheetData) -> int | None:
    for row_index, row in enumerate(sheet.rows):
        fields = {HEADER_ALIASES.get(_normal_header(value)) for value in row if value}
        if {"资产IP", "资产类型", "定级对象名称"}.issubset(fields):
            return row_index
    return None


def _column_schema(sheet: SheetData, header_row: int) -> tuple[dict[int, str], list[str]]:
    width = 0
    for row in sheet.rows[header_row:]:
        for col_index, value in enumerate(row):
            if value:
                width = max(width, col_index + 1)

    mapping: dict[int, str] = {}
    extra_fields: list[str] = []
    used: dict[str, int] = {}
    for col_index in range(width):
        raw_header = sheet.raw(header_row, col_index)
        canonical = HEADER_ALIASES.get(_normal_header(raw_header))
        if canonical:
            mapping[col_index] = canonical
            continue
        base = re.sub(r"[\s\u3000]+", "", raw_header) if raw_header else f"未命名列{col_index + 1}"
        used[base] = used.get(base, 0) + 1
        field = base if used[base] == 1 else f"{base}#{used[base]}"
        mapping[col_index] = field
        extra_fields.append(field)
    return mapping, extra_fields


def parse(path: str | Path) -> dict[str, Any]:
    sheets = read_workbook(path)
    parsed_sheets: list[tuple[SheetData, int, dict[int, str]]] = []
    extra_fields: list[str] = []
    for sheet in sheets:
        header_row = _find_header(sheet)
        if header_row is None:
            continue
        mapping, sheet_extras = _column_schema(sheet, header_row)
        parsed_sheets.append((sheet, header_row, mapping))
        for field in sheet_extras:
            if field not in extra_fields:
                extra_fields.append(field)

    if not parsed_sheets:
        raise ValueError(f"未找到资产清单表头：{path}")

    output_fields = [*ASSET_FIELDS, *extra_fields]
    records: list[dict[str, str]] = []
    for sheet, header_row, mapping in parsed_sheets:
        normalized_header = [_normal_header(sheet.raw(header_row, col_index)) for col_index in mapping]
        for row_index in range(header_row + 1, sheet.row_count):
            values = [sheet.value(row_index, col_index) for col_index in mapping]
            if not any(values):
                continue
            if [_normal_header(value) for value in values] == normalized_header:
                continue
            record = {field: "" for field in output_fields}
            for col_index, field in mapping.items():
                record[field] = sheet.value(row_index, col_index)
            records.append(record)

    object_names: list[str] = []
    for record in records:
        name = record["定级对象名称"]
        if name and name not in object_names:
            object_names.append(name)

    return {
        "定级对象名称": "\n".join(object_names),
        "资产清单": records,
    }


def convert(input_path: str | Path, output_json: str | Path) -> dict[str, Any]:
    result = parse(input_path)
    destination = Path(output_json)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="将资产清单解析为业务字段 JSON")
    parser.add_argument("input", help="输入 .xls/.xlsx 文件")
    parser.add_argument("output", nargs="?", help="输出 JSON；默认与输入文件同名")
    args = parser.parse_args()
    source = Path(args.input)
    destination = Path(args.output) if args.output else source.with_suffix(".json")
    payload = convert(source, destination)
    print(json.dumps({"输出文件": str(destination), "资产数": len(payload["资产清单"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
