from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import xlrd


CHECK_FIELDS = (
    "序号",
    "适用范围",
    "检查类型",
    "检查内容",
    "适用网络属性",
    "检查点",
    "备注",
    "答案",
    "答案说明",
    "分析建议",
)
OUTPUT_CHECK_FIELDS = ("检查分组", *CHECK_FIELDS, "佐证材料", "图片")
DISPIMG_RE = re.compile(r"DISPIMG.*?ID_[A-Za-z0-9]+", re.IGNORECASE)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def visible_text(value: str) -> str:
    return "有图" if DISPIMG_RE.search(value or "") else value


def normal_header(value: str) -> str:
    return re.sub(r"[\s\u3000]+", "", value).replace("：", "").replace(":", "")


HEADER_ALIASES = {normal_header(field): field for field in CHECK_FIELDS}


@dataclass
class SheetData:
    title: str
    rows: list[list[str]]
    merged_ranges: list[tuple[int, int, int, int]]

    def __post_init__(self) -> None:
        self._merged_values: dict[tuple[int, int], str] = {}
        for row_start, row_end, col_start, col_end in self.merged_ranges:
            anchor = self.raw(row_start, col_start)
            for row_index in range(row_start, row_end):
                for col_index in range(col_start, col_end):
                    self._merged_values[(row_index, col_index)] = anchor

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def column_count(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def raw(self, row_index: int, col_index: int) -> str:
        if row_index < 0 or row_index >= len(self.rows):
            return ""
        row = self.rows[row_index]
        if col_index < 0 or col_index >= len(row):
            return ""
        return row[col_index]

    def value(self, row_index: int, col_index: int) -> str:
        raw = self.raw(row_index, col_index)
        if raw:
            return raw
        return self._merged_values.get((row_index, col_index), "")


@dataclass
class ParsedBusiness:
    result: dict[str, Any]
    sheets: list[SheetData]
    row_records: dict[tuple[int, int], dict[str, Any]]
    fields_by_position: dict[tuple[int, int, int], str]


def _read_xlsx(path: Path) -> list[SheetData]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    sheets: list[SheetData] = []
    for worksheet in workbook.worksheets:
        rows = [
            [clean_text(worksheet.cell(row_index, col_index).value) for col_index in range(1, worksheet.max_column + 1)]
            for row_index in range(1, worksheet.max_row + 1)
        ]
        merged = [
            (item.min_row - 1, item.max_row, item.min_col - 1, item.max_col)
            for item in worksheet.merged_cells.ranges
        ]
        sheets.append(SheetData(worksheet.title, rows, merged))
    workbook.close()
    return sheets


def _xls_cell_text(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> str:
    if cell.ctype == xlrd.XL_CELL_DATE:
        return clean_text(xlrd.xldate_as_datetime(cell.value, book.datemode))
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    return clean_text(cell.value)


def _read_xls(path: Path) -> list[SheetData]:
    try:
        workbook = xlrd.open_workbook(path, formatting_info=True)
    except (NotImplementedError, xlrd.XLRDError):
        workbook = xlrd.open_workbook(path)
    sheets: list[SheetData] = []
    for worksheet in workbook.sheets():
        rows = [
            [_xls_cell_text(workbook, worksheet.cell(row_index, col_index)) for col_index in range(worksheet.ncols)]
            for row_index in range(worksheet.nrows)
        ]
        sheets.append(SheetData(worksheet.name, rows, [tuple(item) for item in worksheet.merged_cells]))
    return sheets


def read_workbook(path: str | Path) -> list[SheetData]:
    source = Path(path)
    if source.suffix.lower() == ".xlsx":
        return _read_xlsx(source)
    if source.suffix.lower() == ".xls":
        return _read_xls(source)
    raise ValueError(f"仅支持 .xls/.xlsx：{source}")


def _find_main_header(sheet: SheetData) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(sheet.rows):
        mapping: dict[str, int] = {}
        for col_index, value in enumerate(row):
            canonical = HEADER_ALIASES.get(normal_header(value))
            if canonical and canonical not in mapping:
                mapping[canonical] = col_index
        if {"序号", "检查点", "答案"}.issubset(mapping):
            return row_index, mapping
    return None


def _first_title(sheet: SheetData, before_row: int) -> str:
    for row_index in range(max(0, before_row)):
        for value in sheet.rows[row_index]:
            if value and "填写说明" not in value and "检查项目" not in value:
                return value
    return ""


def _instructions(sheet: SheetData, before_row: int) -> str:
    for row_index in range(max(0, before_row)):
        row = sheet.rows[row_index]
        for col_index, value in enumerate(row):
            if normal_header(value) == "填写说明":
                return "\n".join(item for item in row[col_index + 1 :] if item)
    return ""


def _nonempty_cells(row: Iterable[str]) -> list[tuple[int, str]]:
    return [(index, value) for index, value in enumerate(row) if value]


def _unique_join(values: Iterable[str]) -> str:
    result: list[str] = []
    for value in values:
        shown = visible_text(value)
        if shown and shown not in result:
            result.append(shown)
    return "\n".join(result)


def _unique_headers(values: list[str], width: int) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for col_index in range(width):
        base = values[col_index] if col_index < len(values) and values[col_index] else f"未命名列{col_index + 1}"
        base = re.sub(r"[\s\u3000]+", "", base)
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}#{counts[base]}")
    return result


def parse_sheets(sheets: list[SheetData], source_label: str = "<内存工作表>") -> ParsedBusiness:
    """Parse already-normalized sheets (shared by Excel and PDF backends)."""
    result: dict[str, Any] = {
        "评测表名称": "",
        "填写说明": "",
        "检查项目": [],
        "备份信息": [],
        "其他附表": {},
        "其他图片": [],
    }
    row_records: dict[tuple[int, int], dict[str, Any]] = {}
    fields_by_position: dict[tuple[int, int, int], str] = {}

    for sheet_index, sheet in enumerate(sheets):
        located = _find_main_header(sheet)
        if located is None:
            header_row = next((index for index, row in enumerate(sheet.rows) if any(row)), None)
            if header_row is None:
                continue
            width = 0
            for row in sheet.rows[header_row:]:
                for col_index, value in enumerate(row):
                    if value:
                        width = max(width, col_index + 1)
            if width == 0:
                continue
            headers = _unique_headers(sheet.rows[header_row], width)
            normalized = {normal_header(value) for value in headers}
            key = "备份信息" if {"业务", "备份内容", "备份方式"}.issubset(normalized) else sheet.title
            target = result["备份信息"] if key == "备份信息" else result["其他附表"].setdefault(key, [])
            for row_index in range(header_row + 1, sheet.row_count):
                values = [sheet.value(row_index, col_index) for col_index in range(width)]
                if not any(values):
                    continue
                if [normal_header(value) for value in values] == [normal_header(value) for value in sheet.rows[header_row][:width]]:
                    continue
                record: dict[str, Any] = {
                    headers[col_index]: visible_text(values[col_index]) for col_index in range(width)
                }
                record["图片"] = []
                target.append(record)
                row_records[(sheet_index, row_index + 1)] = record
                for col_index, header in enumerate(headers):
                    fields_by_position[(sheet_index, row_index + 1, col_index + 1)] = header
            continue

        header_row, mapping = located
        if not result["评测表名称"]:
            result["评测表名称"] = _first_title(sheet, header_row)
        if not result["填写说明"]:
            result["填写说明"] = _instructions(sheet, header_row)
        known_columns = set(mapping.values())
        current_group = ""
        for row_index in range(header_row + 1, sheet.row_count):
            raw_row = sheet.rows[row_index]
            raw_nonempty = _nonempty_cells(raw_row)
            if not raw_nonempty:
                continue
            if len(raw_nonempty) == 1 and raw_nonempty[0][0] == mapping["序号"]:
                candidate = raw_nonempty[0][1]
                if not re.fullmatch(r"[0-9A-Za-z]+(?:[-－][0-9A-Za-z]+)+", candidate):
                    current_group = candidate
                    continue

            record: dict[str, Any] = {field: "" for field in OUTPUT_CHECK_FIELDS}
            record["图片"] = []
            record["检查分组"] = current_group
            for field in CHECK_FIELDS:
                if field in mapping:
                    record[field] = visible_text(sheet.value(row_index, mapping[field]))
                    fields_by_position[(sheet_index, row_index + 1, mapping[field] + 1)] = field
            record["佐证材料"] = _unique_join(
                sheet.raw(row_index, col_index)
                for col_index in range(sheet.column_count)
                if col_index not in known_columns
            )
            for col_index in range(sheet.column_count):
                if col_index not in known_columns:
                    fields_by_position[(sheet_index, row_index + 1, col_index + 1)] = "佐证材料"
            result["检查项目"].append(record)
            row_records[(sheet_index, row_index + 1)] = record

    if not result["检查项目"]:
        raise ValueError(f"未找到符合性评测表的检查项目表头：{source_label}")
    return ParsedBusiness(result, sheets, row_records, fields_by_position)


def parse_business(path: str | Path) -> ParsedBusiness:
    return parse_sheets(read_workbook(path), str(path))
